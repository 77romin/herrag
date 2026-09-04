# -*- coding: utf-8 -*-
"""부록 A(개념·지표)를 2장으로 끌어올리고 장 번호와 상호참조를 고친다."""
import re, sys
from docx import Document
sys.stdout.reconfigure(encoding="utf-8")
DOCX = r"C:\SSAFY\chatbot-project_lab\ChatBot_고찰보고서.docx"
doc = Document(DOCX)
body = doc.element.body

def find_p(pred):
    for p in doc.paragraphs:
        if pred(p):
            return p._p
    return None

app_h = find_p(lambda p: p.text.strip().startswith("부록 A."))
tgt_h = find_p(lambda p: p.text.strip().startswith("2. F207 ①"))
assert app_h is not None and tgt_h is not None

kids = list(body)
ai, ti = kids.index(app_h), kids.index(tgt_h)
# 부록 앞의 쪽나눔 문단도 함께 옮긴다
start = ai - 1 if ai > 0 and kids[ai - 1].tag.endswith("}p") and not kids[ai - 1].xpath(".//w:t") else ai
end = len(kids)
while end > start and kids[end - 1].tag.endswith("}sectPr"):
    end -= 1
block = kids[start:end]
print(f"이동 블록 {len(block)}개 (부록 A ~ 문서 끝)")
for el in block:
    tgt_h.addprevious(el)

# ── 장 번호 재배치: 기존 2~11 -> 3~12, 부록 A -> 2 ──────────────
def retext(par, new):
    if par.runs:
        par.runs[0].text = new
        for r in par.runs[1:]:
            r.text = ""

for par in doc.paragraphs:
    if not par.style.name.startswith("Heading"):
        continue
    t = par.text.strip()
    m = re.match(r"^(\d+)([.\-])", t)
    if m and 2 <= int(m.group(1)) <= 11:
        retext(par, re.sub(r"^\d+", str(int(m.group(1)) + 1), t, count=1))

for par in doc.paragraphs:
    t = par.text.strip()
    if t.startswith("부록 A."):
        retext(par, t.replace("부록 A.", "2.", 1))
    elif re.match(r"^A-\d\.", t):
        retext(par, "2-" + t[2:])

# ── 상호참조 ────────────────────────────────────────────────
FIX = [
    ("검색 방식과 평가 지표의 뜻은 부록 A에 따로 정리했다.",
     "검색 방식과 평가 지표의 뜻은 다음 장에 정리했다."),
    ("본문에 나오는 검색 방식과 평가 지표를 정리한다.",
     "이 보고서에서 쓰는 검색 방식과 평가 지표를 먼저 정리한다."),
    ("3장에서 측정한 대로", "4장에서 측정한 대로"),
    ("7-2절에서 확인했듯", "8-2절에서 확인했듯"),
    ("5장에서 확인했듯", "6장에서 확인했듯"),
    ("본문 4-1절의 명찰 사례가 여기에 해당한다.",
     "본문 5-1절의 명찰 사례가 여기에 해당한다."),
    ("본문 4-5절에서 확인한 대응이 바로 이 구조에서 나온다.",
     "본문 5-5절의 대응이 바로 이 구조에서 나온다."),
    ("본문 7-1절에서 점수의 절반이 음수였다고 밝혔다.",
     "본문 8-1절에서 다루듯 점수의 절반이 음수였다."),
]
done = 0
for par in doc.paragraphs:
    t = par.text
    for a, b in FIX:
        if a in t:
            retext(par, t.replace(a, b)); done += 1; break
print(f"상호참조 {done}/{len(FIX)}곳 수정")

doc.save(DOCX)
print("저장 완료")
